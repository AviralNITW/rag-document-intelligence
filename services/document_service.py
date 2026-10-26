import os
import uuid
import logging
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional

# Must precede `import chromadb`: it reads ANONYMIZED_TELEMETRY at import time,
# and its bundled posthog client is API-incompatible (every capture() call fails
# with a TypeError). This is a local app, so telemetry is disabled at the source
# instead of spamming the server log on every collection operation.
os.environ.setdefault("ANONYMIZED_TELEMETRY", "False")

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def _disable_chroma_telemetry() -> None:
    """Neutralizes chromadb's posthog telemetry client."""
    try:
        from chromadb.telemetry.product.posthog import Posthog

        Posthog.capture = lambda self, event: None
        Posthog._direct_capture = lambda self, event: None
    except Exception as e:
        logger.debug("Could not neutralize chromadb telemetry: %s", e)

MAX_FILE_SIZE_BYTES = 200 * 1024 * 1024  # 200 MB maximum
CHROMA_DIR = os.path.join(".", "chroma_db")

# UI format-filter label -> stored document "type" (extension upper-cased).
FORMAT_FILTER_MAP: Dict[str, Optional[List[str]]] = {
    "All Formats": None,
    "PDF": ["PDF"],
    "DOCX": ["DOCX"],
    "TXT": ["TXT"],
    "Markdown": ["MD"],
}


def format_file_size(size_bytes: int) -> str:
    """Formats bytes into human-readable size."""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"


def relative_time(timestamp_str: str, fmt: str = "%Y-%m-%d %H:%M:%S") -> str:
    """Converts a stored timestamp string into a live relative label."""
    try:
        dt = datetime.strptime(timestamp_str, fmt)
    except (TypeError, ValueError):
        return "—"
    seconds = (datetime.now() - dt).total_seconds()
    if seconds < 60:
        return "Just now"
    if seconds < 3600:
        return f"{int(seconds // 60)}m ago"
    if seconds < 86400:
        return f"{int(seconds // 3600)}h ago"
    return f"{int(seconds // 86400)}d ago"


def filter_documents_by_format(documents: List[Dict[str, Any]], format_label: str) -> List[Dict[str, Any]]:
    """Filters document records by UI format label using FORMAT_FILTER_MAP.

    An unknown label returns an empty list (fail loud) rather than silently
    showing every document as if no filter were applied.
    """
    if format_label not in FORMAT_FILTER_MAP:
        return []
    allowed = FORMAT_FILTER_MAP[format_label]
    if allowed is None:
        return list(documents)
    return [d for d in documents if d.get("type") in allowed]


def collection_name_for(session_id: str) -> str:
    return f"session_{session_id.replace('-', '_')}"


def _chroma_client() -> Any:
    _disable_chroma_telemetry()
    import chromadb
    return chromadb.PersistentClient(path=CHROMA_DIR)


def extract_document_pages(file_path: str, original_filename: str) -> List[Document]:
    """Extracts pages and text content from supported file types.

    Document text is stored verbatim: it is untrusted DATA that the RAG prompt
    explicitly fences off, rather than something we rewrite at ingest time.
    """
    ext = os.path.splitext(original_filename)[1].lower()
    docs: List[Document] = []

    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            table_str = "\n".join([
                                " | ".join([cell.replace("\n", " ") if cell else "" for cell in row])
                                for row in table if any(row)
                            ])
                            if table_str.strip():
                                page_text += "\n\nTable:\n" + table_str

                    if page_text.strip():
                        docs.append(Document(
                            page_content=page_text,
                            metadata={
                                "source": original_filename,
                                "page": page_idx,
                                "file_type": "PDF"
                            }
                        ))
        except Exception as e:
            logger.warning("pdfplumber extraction failed (%s); falling back to PyPDFLoader.", e)

        if not docs:
            from langchain_community.document_loaders import PyPDFLoader
            loader = PyPDFLoader(file_path)
            for doc in loader.load():
                if doc.page_content.strip():
                    doc.metadata["source"] = original_filename
                    doc.metadata["file_type"] = "PDF"
                    docs.append(doc)

    elif ext == ".docx":
        import docx
        doc_obj = docx.Document(file_path)
        full_text = []
        for para in doc_obj.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc_obj.tables:
            for row in table.rows:
                row_str = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_str:
                    full_text.append(row_str)
        combined = "\n\n".join(full_text)
        if combined.strip():
            docs.append(Document(
                page_content=combined,
                # Non-PDF formats have no meaningful page number: the key is omitted
                # (Chroma rejects None metadata values) and readers treat an absent
                # page as "Whole document".
                metadata={
                    "source": original_filename,
                    "file_type": "DOCX"
                }
            ))

    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        if content.strip():
            file_type = "Markdown" if ext == ".md" else "Plain Text"
            docs.append(Document(
                page_content=content,
                metadata={
                    "source": original_filename,
                    "file_type": file_type
                }
            ))
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Supported: .pdf, .docx, .txt, .md")

    if not docs or not any(d.page_content.strip() for d in docs):
        raise ValueError(f"Could not extract readable text from '{original_filename}'.")

    return docs


def process_and_index_document(
    file_path: str,
    original_filename: str,
    session_id: str,
    embeddings_model
) -> Tuple[Dict[str, Any], List[Document]]:
    """
    Parses, splits, embeds, and stores a document into ChromaDB.
    Returns metadata record and chunk documents.
    """
    file_size_bytes = os.path.getsize(file_path)
    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        raise ValueError("File exceeds maximum allowed limit of 200 MB.")

    raw_pages = extract_document_pages(file_path, original_filename)

    from langchain_text_splitters import RecursiveCharacterTextSplitter
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150,
        separators=["\n\n", "\n", " ", ""]
    )
    chunks = splitter.split_documents(raw_pages)

    # Enrich chunks with index IDs
    for idx, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = idx
        chunk.metadata["total_chunks"] = len(chunks)
        chunk.metadata["session_id"] = session_id
        # Chroma only accepts str/int/float/bool metadata values, so any loader
        # that yields None (e.g. a missing page number) is stripped here rather
        # than failing the whole upload with a metadata validation error.
        chunk.metadata = {k: v for k, v in chunk.metadata.items() if v is not None}

    # Store into ChromaDB
    from langchain_community.vectorstores import Chroma
    collection = collection_name_for(session_id)
    Chroma.from_documents(
        documents=chunks,
        embedding=embeddings_model,
        collection_name=collection,
        persist_directory=CHROMA_DIR
    )

    doc_record = {
        "id": str(uuid.uuid4()),
        "name": original_filename,
        "type": os.path.splitext(original_filename)[1].upper().replace(".", ""),
        "size_bytes": file_size_bytes,
        "size_str": format_file_size(file_size_bytes),
        "chunks_count": len(chunks),
        "status": "Indexed",
        "upload_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "time_ago": "Just now",
        "page_count": len(raw_pages)
    }

    return doc_record, chunks


def compute_insights(documents: List[Dict[str, Any]], all_chunks: List[Any]) -> Dict[str, Any]:
    """Computes Document Insights from actual repository state (live timestamps)."""
    total_chunks = len(all_chunks)
    total_bytes = sum(doc.get("size_bytes", 0) for doc in documents)
    total_size_str = format_file_size(total_bytes) if total_bytes > 0 else "0 KB"
    doc_count = len(documents)
    last_indexed = relative_time(documents[-1]["upload_time"]) if documents and documents[-1].get("upload_time") else "—"

    return {
        "total_chunks": total_chunks,
        "file_size": total_size_str,
        "doc_count": doc_count,
        "last_indexed": last_indexed
    }


# ── Vector store maintenance ─────────────────────────────
def clear_session_storage(session_id: str) -> bool:
    """Purges the session-isolated vector collection. Returns True if deleted."""
    name = collection_name_for(session_id)
    try:
        _chroma_client().delete_collection(name)
        logger.info("Deleted vector collection '%s'.", name)
        return True
    except Exception as e:
        logger.info("Collection cleanup notice for '%s': %s", name, e)
        return False


def delete_document_chunks(session_id: str, source_name: str) -> int:
    """Deletes all chunks for one source document from the session collection."""
    name = collection_name_for(session_id)
    try:
        client = _chroma_client()
        collection = client.get_collection(name)
        result = collection.get(where={"source": source_name}, include=[])
        ids = result.get("ids", [])
        if ids:
            collection.delete(ids=ids)
        return len(ids)
    except Exception as e:
        logger.info("Chunk deletion notice for '%s' in '%s': %s", source_name, name, e)
        return 0


def list_session_collections() -> List[str]:
    """Lists all vector store collections on disk (without creating any)."""
    try:
        collections = _chroma_client().list_collections()
        names = []
        for c in collections:
            names.append(c if isinstance(c, str) else getattr(c, "name", str(c)))
        return names
    except Exception as e:
        logger.warning("Could not list vector collections: %s", e)
        return []


def purge_all_vector_stores() -> int:
    """Deletes every collection on disk (all sessions). Returns count deleted."""
    deleted = 0
    for name in list_session_collections():
        try:
            _chroma_client().delete_collection(name)
            deleted += 1
        except Exception as e:
            logger.warning("Could not delete collection '%s': %s", name, e)
    return deleted